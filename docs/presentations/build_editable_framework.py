"""Build a native-shape, editable PowerPoint. Requires python-pptx and Pillow."""
from pathlib import Path
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/presentations'
R=Presentation();R.slide_width=Inches(30);R.slide_height=Inches(20)
S=R.slides.add_slide(R.slide_layouts[6]); unit=30/1600
C={'ink':'162D43','line':'41566A','blue':'4B8FDA','blue2':'C2DCF9','teal':'2DADAE','teal2':'D9F1EE','purple':'9262C9','purple2':'F0E9F9','orange':'E99A48','orange2':'FFF0DD','white':'FFFFFF','gray':'F3F6F9'}
def pos(v):return Inches(v*unit)
def col(c):return RGBColor.from_string(C.get(c,c))
def shape(x,y,w,h,kind=MSO_SHAPE.RECTANGLE,fill='white',stroke='line',width=1,slide=None):
 o=(slide or S).shapes.add_shape(kind,pos(x),pos(y),pos(w),pos(h));o.fill.solid();o.fill.fore_color.rgb=col(fill)
 if stroke:o.line.color.rgb=col(stroke);o.line.width=Pt(width)
 else:o.line.fill.background()
 # Avoid theme-dependent drop shadows and overly rounded corners.
 if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
  o.adjustments[0] = .08
 o._element.spPr.append(OxmlElement('a:effectLst'))
 return o
def txt(x,y,w,h,t,size=15,bold=False,color='ink',align='center',slide=None,font='Arial'):
 o=(slide or S).shapes.add_textbox(pos(x),pos(y),pos(w),pos(h));tf=o.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=pos(2);tf.margin_top=tf.margin_bottom=0
 tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for i,l in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[align]
  p.font.name=font;p.font.size=Pt(size*1.32);p.font.bold=bold;p.font.color.rgb=col(color);p.space_before=p.space_after=0
 return o
def box(x,y,w,h,t,fill='gray',stroke='line',size=14,bold=False):
 o=shape(x,y,w,h,MSO_SHAPE.ROUNDED_RECTANGLE,fill,stroke);txt(x+3,y+2,w-6,h-4,t,size,bold);return o
def line(points,color='line',dashed=False,arrow=True,slide=None):
 for i,(a,b) in enumerate(zip(points,points[1:])):
  o=(slide or S).shapes.add_connector(MSO_CONNECTOR.STRAIGHT,pos(a[0]),pos(a[1]),pos(b[0]),pos(b[1]));o.line.color.rgb=col(color);o.line.width=Pt(1.3)
  ln=o._element.spPr.get_or_add_ln()
  if dashed:
   e=OxmlElement('a:prstDash');e.set('val','dash');ln.append(e)
  if arrow and i==len(points)-2:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
 return o
def tensor(x,y,w=42,h=28,fill='blue',depth=9):
 g=S.shapes.add_group_shape()
 for k in [2,1,0]:
  o=g.shapes.add_shape(MSO_SHAPE.RECTANGLE,pos(x+k*depth/2),pos(y-k*depth/2),pos(w),pos(h));o.fill.solid();o.fill.fore_color.rgb=col(fill if k==0 else ('blue2' if fill=='blue' else 'teal2'));o.line.color.rgb=col('white');o.line.width=Pt(.6)
 return g
def grid(x,y,w=70,h=55,n=5,palette=None):
 g=S.shapes.add_group_shape();pa=palette or ['D5F0EC','8BD3D0','48B5BB','2D929E']
 for j in range(n):
  for i in range(n):
   o=g.shapes.add_shape(MSO_SHAPE.RECTANGLE,pos(x+i*w/n),pos(y+j*h/n),pos(w/n-.6),pos(h/n-.6));o.fill.solid();o.fill.fore_color.rgb=col(pa[(i*2+j*3)%len(pa)]);o.line.fill.background()
 return g
def encoder(x,y,w,h,label,fill):
 shape(x,y,w,h,MSO_SHAPE.TRAPEZOID,fill,'line');txt(x+3,y+5,w-6,h-10,label,16,True)
def mlp(x,y,label,fill='blue',w=24,h=34):
 shape(x,y,w,h,MSO_SHAPE.TRAPEZOID,fill,'line');txt(x-3,y+2,w+6,h-4,label,13,True)
def circ(x,y,label,fill='white',d=25):
 shape(x,y,d,d,MSO_SHAPE.OVAL,fill,'line');txt(x,y,d,d,label,17,True)
def tokens(x,y,count=4):
 for i in range(count):shape(x+i*16,y,11,17,MSO_SHAPE.ROUNDED_RECTANGLE,'purple2','purple')
def panel(x,y,w,h,label,stroke,fill='white'):
 shape(x,y,w,h,MSO_SHAPE.ROUNDED_RECTANGLE,fill,stroke,1.1);txt(x+10,y+3,w-20,26,label,19,True,stroke,'left',font='Georgia')
# Paper-style page, clear hierarchy and native visual objects.
txt(15,5,1555,32,'(a) Quality-aware multimodal pretraining',23,True,align='left',font='Georgia')
# Input thumbnails are entirely native vector scenery.
for y,label in [(58,'SafeAug 1'),(153,'SafeAug 2'),(248,'Clean')]:
 txt(6,y+20,88,30,label,14,True)
 shape(102,y,80,63,fill='blue2',stroke='blue');shape(104,y+35,76,25,fill='teal2',stroke=None)
 shape(113,y+16,27,27,MSO_SHAPE.ISOSCELES_TRIANGLE,'blue',None);shape(140,y+24,32,22,MSO_SHAPE.ISOSCELES_TRIANGLE,'teal',None)
 shape(126,y+47,26,7,MSO_SHAPE.TRAPEZOID,'orange',None)
 line([(182,y+31),(199,y+31),(199,181),(215,181)])
panel(215,53,265,276,'Shared visual backbone','line')
encoder(233,98,91,77,'CLIP\nViT-L/14','blue2');tensor(344,121,58,20);txt(326,91,130,24,'Global · 768-D',12)
encoder(233,207,91,78,'ResNet-50\nlayer4','teal2');grid(345,223,53,48);txt(321,279,142,23,'7 × 7 × 2048',12)
box(421,159,46,70,'F\nu\ns\ni\no\nn','gray',size=11)
line([(324,137),(339,137)]);line([(402,132),(413,132),(413,176),(421,176)])
line([(324,245),(340,245)]);line([(399,245),(413,245),(413,212),(421,212)])
line([(467,195),(496,195)]);tensor(500,180,18,44);txt(490,145,44,25,'z',15,True)
# Quality projection, all PQ shared, predictors SafeAug only.
txt(542,40,385,25,'Shared P_Q / D_Q across augmented views',12,True,'purple')
for y,label in [(78,'z₁'),(172,'z₂'),(266,'z_c')]:
 txt(535,y,37,29,label,14);mlp(582,y,'P_Q','blue2',42,38);line([(572,y+19),(582,y+19)]);tensor(644,y+9,50,15);line([(624,y+19),(642,y+19)])
 if label!='z_c':
  mlp(719,y,'D_Q','orange2',42,38);line([(696,y+19),(719,y+19)]);tensor(779,y+8,32,17);line([(761,y+19),(777,y+19)])
 else:
  box(727,y-9,243,60,'Component anchors + ordinal\nFive-level quality text CE','purple2','purple',13)
  line([(695,y+19),(727,y+19)])
# Feature bus explicitly z1/z2/zc, not cross-stage sharing.
line([(520,200),(529,200),(529,97),(535,97)]);line([(529,200),(535,191)]);line([(529,200),(529,285),(535,285)])
box(837,90,133,120,'Symmetric\nSimSiam\n\nsg targets','gray',size=14)
line([(812,95),(837,119)]);line([(812,189),(837,176)])
box(582,351,178,48,'Separate MOS heads','orange2','orange',13,True)
line([(520,210),(534,210),(534,375),(582,375)])
txt(767,351,203,48,'Quality / Artifact / Aesthetics\nClean backbone supervision',12,align='left')
# Compact independent semantic relation panel.
panel(1005,52,571,277,'Separate image–text relation space','purple')
tensor(1026,106,45,33);mlp(1100,103,'P_I','blue2',40,40);mlp(1175,103,'D_I','orange2',40,40)
line([(1072,122),(1100,122)]);line([(1140,122),(1175,122)])
tokens(1026,192);encoder(1100,174,106,62,'Text encoder\nFrozen','purple2');mlp(1240,183,'P_T','purple2',40,40);mlp(1310,183,'D_T','orange2',40,40)
line([(1088,201),(1100,201)]);line([(1206,201),(1240,201)]);line([(1280,201),(1310,201)])
txt(1018,246,348,48,'Prompt / Answer / region phrases\nTrainable contexts; relation-specific adapters',12)
grid(1411,124,103,90,6,['F0ECFA','D3C7EE','B5A7E0','7C70BA','BCD9F2'])
line([(1215,123),(1369,123),(1369,145),(1406,145)]);line([(1350,203),(1406,203)])
txt(1380,223,168,57,'Image–Prompt / Answer\nPrompt–Answer\nRegion–word',12)
box(1016,348,555,57,'P_Q → Quality Adapter initialization\nBackbone / contexts / alignment adapters / quality bank transferred','purple2','purple',13)
line([(24,427),(1576,427)],dashed=True,arrow=False)
# Downstream shared feature extraction
# key named feature ports avoid long crossovers.
txt(15,439,1555,32,'(b) Downstream quality & alignment assessment',23,True,align='left',font='Georgia')
box(21,567,98,90,'Input image\n224 × 224','blue2','blue',14)
encoder(148,498,111,89,'CLIP\nViT-L/14','blue2');tensor(285,533,48,19);txt(269,504,87,20,'768-D',12)
encoder(148,688,111,89,'ResNet-50\nlayer3','teal2');grid(285,708,56,49);txt(259,773,125,23,'14 × 14 × 1024',11)
line([(119,611),(133,611),(133,542),(148,542)]);line([(133,611),(133,732),(148,732)])
box(365,521,62,44,'768→256','blue2','blue',11);line([(335,542),(365,542)])
box(365,710,62,44,'1024→256','teal2','teal',10);line([(344,732),(365,732)])
panel(454,496,217,323,'Semantic–Structure Fusion','line')
tensor(476,551,70,18);txt(471,526,114,23,'Global query Q',12)
grid(476,697,65,60);grid(568,697,65,60);txt(474,764,161,24,'Local keys / values',12)
grid(529,599,93,69,6,['E5EFF9','AACBE8','8EADD4','DDDDDC','E7CA97'])
line([(546,560),(556,560),(556,594)]);line([(506,697),(506,636),(525,636)]);line([(600,697),(600,673)])
line([(427,543),(447,543),(447,560),(472,560)]);line([(427,732),(449,732),(449,726),(472,726)])
txt(474,791,177,24,'z: 256-D · shared by Q / A',12,True)
box(281,833,149,50,'Shared local projection L\nP₀: 196 × 256','teal2','teal',11)
line([(396,754),(396,815),(356,815),(356,833)])
# QUALITY branch
panel(699,493,877,281,'Quality branch (Q)','teal','F1FAF8')
tensor(715,553,16,35);txt(705,525,36,20,'z',12,True)
panel(758,533,184,70,'Quality Adapter','teal')
for x in [775,830,885]:mlp(x,566,'MLP','teal2',34,24)
line([(735,571),(758,571)]);line([(808,578),(828,578)]);line([(864,578),(884,578)])
tensor(955,560,24,15,fill='teal');txt(951,535,34,20,'a',13,True)
line([(942,578),(953,578)]);circ(1001,582,'+')
line([(979,568),(1013,568),(1013,582)])
tensor(763,620,21,15);txt(744,638,68,22,'z bypass',10)
line([(784,627),(984,627),(984,595),(1001,595)])
box(1055,548,128,48,'Global head','gray',size=14)
line([(1026,595),(1042,595),(1042,572),(1055,572)])
tensor(1204,560,16,29);txt(1188,531,57,24,'Q_g',14,True)
line([(1183,572),(1203,572)])
txt(798,659,426,29,'r = z + α · sg(‖z‖) · normalize(a)',14,font='Cambria')
# Quality local is one grouped block with correct parallel visual scores/weights.
tensor(720,704,28,27,fill='teal');txt(708,681,50,21,'P₀',12,True)
panel(774,695,450,67,'NSS-guided local quality','teal')
# keep information in compact parallel text rows rather than incorrect serial NSS paths
box(788,728,113,24,'Pq = P₀ + N(NSS)','teal2','teal',10)
grid(921,733,71,13,6,['F7DB9E','C7E8E2','9FC4E9','D0BBE0']);grid(1024,733,70,13,6,['DDE7F5','D1B4DC','EACB8C','8FBCE3'])
txt(912,716,88,17,'Patch scores',9);txt(1015,716,88,17,'NSS weights',9)
circ(1113,725,'Σ',d=24);line([(1095,740),(1113,740)]);line([(992,740),(1104,740),(1104,732),(1113,732)]);line([(1095,740),(1107,740),(1107,747),(1113,747)])
txt(1143,726,60,25,'Q_l',13,True);line([(1138,738),(1145,738)])
line([(748,718),(761,718),(761,739),(788,739)])
# Fusion separate actual inputs
circ(1324,634,'β',d=29);line([(1220,575),(1276,575),(1276,642),(1324,642)])
line([(1224,738),(1276,738),(1276,656),(1324,656)])
box(1405,627,147,60,'Perceptual\nQuality Q','teal2','teal',17,True);line([(1353,648),(1405,648)])
txt(1279,689,285,28,'Q = (1−β)Q_g + βQ_l',14,font='Cambria');txt(1308,719,226,22,'β: learned scalar',11)
# Auxiliary from a only, training-only separate note.
box(1101,490,225,39,'Five-level text CE · auxiliary\nλ=0.005 · score blend=0','purple2','purple',10)
line([(966,558),(966,516),(1101,516)],color='purple',dashed=True)
# ALIGNMENT branch
panel(699,796,877,238,'Alignment branch (A)','orange','FFFAF3')
# global image path
texty=864
tensor(715,853,16,29);txt(705,831,36,20,'z',12,True)
box(758,846,138,46,'Image adapter A_G','orange2','orange',12);line([(734,866),(758,866)])
tensor(918,858,16,21,fill='blue');txt(914,837,30,20,'v',12,True)
line([(896,869),(918,869)])
box(1020,834,175,52,'Global pair head\n[v, u, |v−u|]','gray',size=12)
line([(936,869),(975,869),(975,850),(1020,850)])
# Prompt input is outside branch but explicitly goes U then AT
box(16,935,148,58,'Prompt + learnable\ncontext tokens','purple2','purple',12)
tokens(183,956);encoder(266,927,116,65,'Frozen CLIP\ntext encoder','purple2')
line([(164,964),(180,964)]);line([(246,964),(266,964)])
box(406,943,54,37,'U','purple2','purple',14);line([(382,961),(406,961)])
box(491,943,146,37,'Text adapter A_T','purple2','purple',12);line([(460,961),(491,961)])
tensor(662,951,17,25,fill='blue');txt(651,979,36,20,'u',12,True);line([(637,961),(662,961)])
# Same u reused named port at global head, not confused with P0.
tensor(971,893,17,18);txt(952,907,53,20,'u',11,True)
line([(988,901),(1005,901),(1005,873),(1020,873)])
tensor(718,958,23,25,fill='teal');txt(707,935,50,20,'P₀',12,True)
panel(773,924,424,91,'Local–text matching','orange')
grid(791,960,54,38);tensor(858,992,17,16);txt(851,1010,25,14,'u',9)
grid(905,960,65,38,5,['E8E4F7','CBB3DE','93B6DB','ECD09F'])
line([(845,973),(901,973)]);line([(876,1000),(892,1000),(892,985),(901,985)])
box(999,958,66,41,'Attention\npooling','orange2','orange',10);line([(972,978),(999,978)])
box(1092,958,91,41,'Local pair head','gray',size=10);line([(1065,978),(1092,978)])
line([(742,971),(773,971)])
txt(914,1000,275,14,'A_P · X_I / X_T · [e, u, |e−u|]',8)
# u enters local group independently as named port. No false text-to-image-grid edge.
line([(680,963),(692,963),(692,1024),(873,1024),(873,1008)],color='purple')
# global and local predictions into sample-conditioned alignment gate
circ(1330,917,'γ',d=30)
line([(1195,859),(1266,859),(1266,925),(1330,925)])
line([(1197,978),(1266,978),(1266,939),(1330,939)])
txt(1203,835,52,20,'A_g',12,True);txt(1203,983,52,20,'A_l',12,True)
box(1405,904,147,61,'Text–Image\nAlignment A','orange2','orange',16,True);line([(1360,932),(1405,932)])
txt(1285,970,280,26,'A = γ A_g + (1−γ) A_l',14,font='Cambria');txt(1301,1000,252,20,'γ: sample-wise, conditioned on v, u',10)
# page footnotes
txt(20,1039,1550,23,'Same z / P₀ / u ports denote the same feature tensors.   sg: stop-gradient.   Visual encoders trainable; text backbone frozen.   Source layer4 → downstream layer3: local projection reset.',10,align='left')
S.notes_slide.notes_text_frame.text='Editable reconstruction of the paper-style framework. All diagram text, tensors, heatmaps, modules and connectors are native PowerPoint objects. Reused named ports avoid crossed lines. Representation transfer resets task heads and downstream local projection. The matching group is schematic: XI/XT similarities pool AP(P0); u is not a local image token. NSS is computed from the image, not from P0. Detailed source: docs/figures/model_forward_and_sharing.md.'
# Reference page, explicitly a raster, so editable page cannot be confused with a screenshot.
ref=R.slides.add_slide(R.slide_layouts[6]);ref.shapes.add_picture(str(ROOT/'docs/figures/paper_framework_v2.png'),0,0,width=R.slide_width,height=R.slide_height)
ref.notes_slide.notes_text_frame.text='Raster visual reference only. This original generated design has known connectivity errors. Use slide 1 for editable corrections; do not use this raster as the exact implementation diagram.'
# Editing guide
G=R.slides.add_slide(R.slide_layouts[6]);txt(70,60,1450,70,'编辑说明 / Editing guide',34,True,slide=G,align='left')
items=[
'第 1 页：可编辑重绘版。文字、公式、模块、箭头、特征张量和热图均为原生 PPT 对象。',
'第 2 页：原始论文风格图，仅作视觉参考。PNG 本身不能在 PowerPoint 中拆成原生图层。',
'在第 1 页直接点击文字修改；选中连接线后可拖动端点。图形填充、线条和字体可直接修改。',
'张量与热图以局部组合保存：右键 → 组合 → 取消组合，可编辑每个小格；也可 Ctrl+单击选择组内对象。',
'全选多个对象后可组合移动。重复的 z / P₀ / u 是同一个特征的命名端口，用于避免长线交叉。',
'已按代码修正关键关系：z→全局 Adapter，P₀→局部匹配；质量标量 β，一致性逐样本 γ。',
'保留原图配色和双面板视觉语言，但原生重绘并非与原 PNG 逐像素一致。',
'导出论文图：文件 → 导出 → PDF；或另存图片。排版时建议保留矢量 PDF。',
'依据：代码 e127834；CLIP Large、下游 layer3、Residual Quality Adapter、CE=0.005、标量质量融合。',
'未纳入实验中的逐图质量门控，也未把学习率消融宣称为已确定主方法。'
]
for i,t in enumerate(items):txt(75,165+i*76,1435,65,t,23,slide=G,align='left')
OUT.mkdir(parents=True,exist_ok=True);R.save(OUT/'paper_framework_editable.pptx')
print('Saved',OUT/'paper_framework_editable.pptx')
print('Native diagram objects:',len(S.shapes),'including',sum(len(g.shapes) for g in S.shapes if hasattr(g,'shapes')),'group children')
